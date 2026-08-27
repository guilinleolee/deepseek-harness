"""test_47_1_docx · 89-financial-analyst V11.0 .docx 财务研报产出验证

验证：
  1. 文件存在 + 大小 > 5000
  2. 至少 3 个 heading（一、二、三）
  3. 至少 1 个表格（财务比率）
  4. 包含 "财务分析简报" 标题
"""
import sys
from datetime import date
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

from office_producers import build_docx_report


def test_47_1_docx():
    out = Path(__file__).resolve().parent.parent / "output" / "docx" / f"test-{date.today().isoformat()}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    result = build_docx_report(out)

    assert out.exists(), f"[FAIL] docx not created: {out}"
    assert out.stat().st_size > 5000, f"[FAIL] docx too small: {out.stat().st_size}"

    doc = Document(out)
    paragraphs = doc.paragraphs
    full_text = "\n".join(p.text for p in paragraphs)
    assert "财务分析简报" in full_text, "[FAIL] 标题 '财务分析简报' not found"
    assert "财务表现" in full_text, "[FAIL] section '财务表现' missing"
    assert "经营分析" in full_text, "[FAIL] section '经营分析' missing"
    assert "投资建议" in full_text, "[FAIL] section '投资建议' missing"

    # 至少 1 个表格（财务比率）
    assert len(doc.tables) >= 1, f"[FAIL] expected ≥1 table, got {len(doc.tables)}"
    # 表格至少 5 行（表头 + 4 指标）
    assert len(doc.tables[0].rows) >= 5, f"[FAIL] table rows {len(doc.tables[0].rows)} < 5"

    print(f"[PASS] docx: {result}")


if __name__ == "__main__":
    test_47_1_docx()
