#!/usr/bin/env python3
"""
MD-to-PPTX Converter

将 Markdown 文件转换为 PowerPoint 演示文稿
"""

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Optional


class LayoutType(Enum):
    """幻灯片布局类型"""
    COVER = "cover"           # 封面
    TITLE = "title"           # 章节标题
    CONTENT = "content"       # 内容页
    TWO_COLUMN = "two-column" # 双栏
    CODE = "code"            # 代码页
    CHART = "chart"          # 图表页
    TABLE = "table"          # 表格页
    ENDING = "ending"        # 结尾页


class SlideStyle(Enum):
    """演示风格"""
    CORPORATE = "corporate"   # 企业蓝
    MINIMAL = "minimal"       # 极简白
    ACADEMIC = "academic"     # 学术灰
    TECH = "tech"            # 科技蓝
    DARK = "dark"            # 暗夜黑
    CREATIVE = "creative"     # 创意彩


@dataclass
class SlideContent:
    """幻灯片内容"""
    layout: LayoutType = LayoutType.CONTENT
    title: str = ""
    subtitle: str = ""
    bullets: list = field(default_factory=list)
    code: str = ""
    code_lang: str = ""
    table_data: list = field(default_factory=list)
    image_url: str = ""
    notes: str = ""
    transition: str = "fade"
    background: Optional[str] = None
    raw_content: list = field(default_factory=list)


@dataclass
class SlideStyleConfig:
    """样式配置"""
    primary_color: str = "#1E3A5F"
    secondary_color: str = "#2D5A87"
    background_color: str = "#FFFFFF"
    text_color: str = "#1A202C"
    accent_color: str = "#3182CE"
    font_heading: str = "Microsoft YaHei"
    font_body: str = "Microsoft YaHei"
    font_code: str = "Consolas"

    @classmethod
    def from_style(cls, style: SlideStyle) -> "SlideStyleConfig":
        """从风格名称创建配置"""
        configs = {
            SlideStyle.CORPORATE: cls(
                primary_color="#1E3A5F",
                secondary_color="#2D5A87",
                background_color="#FFFFFF",
                text_color="#1A202C",
                accent_color="#3182CE"
            ),
            SlideStyle.MINIMAL: cls(
                primary_color="#1A202C",
                secondary_color="#4A5568",
                background_color="#FFFFFF",
                text_color="#1A202C",
                accent_color="#718096"
            ),
            SlideStyle.ACADEMIC: cls(
                primary_color="#4A5568",
                secondary_color="#718096",
                background_color="#F7FAFC",
                text_color="#2D3748",
                accent_color="#805AD5"
            ),
            SlideStyle.TECH: cls(
                primary_color="#0D1B2A",
                secondary_color="#1B3A4B",
                background_color="#0D1B2A",
                text_color="#E0E7FF",
                accent_color="#60A5FA"
            ),
            SlideStyle.DARK: cls(
                primary_color="#1A202C",
                secondary_color="#2D3748",
                background_color="#1A202C",
                text_color="#F7FAFC",
                accent_color="#4FD1C5"
            ),
            SlideStyle.CREATIVE: cls(
                primary_color="#805AD5",
                secondary_color="#9F7AEA",
                background_color="#FAF5FF",
                text_color="#2D3748",
                accent_color="#ED64A6"
            ),
        }
        return configs.get(style, configs[SlideStyle.CORPORATE])


@dataclass
class ConversionResult:
    """转换结果"""
    success: bool
    slides: list[SlideContent] = field(default_factory=list)
    output_path: Optional[Path] = None
    error_message: Optional[str] = None
    stats: dict = field(default_factory=dict)


class MarkdownParser:
    """Markdown 解析器"""

    def __init__(self, content: str):
        self.content = content
        self.lines = content.split('\n')
        self.current_index = 0

    def parse(self) -> list[SlideContent]:
        """解析 Markdown 内容为幻灯片列表"""
        slides = []
        current_slide = SlideContent()

        for line_num, line in enumerate(self.lines, 1):
            # 检查幻灯片分隔符
            if self._is_slide_separator(line):
                if current_slide.title or current_slide.bullets:
                    slides.append(current_slide)
                current_slide = SlideContent()
                continue

            # 检查注释指令
            if self._is_comment(line):
                self._process_comment(line, current_slide)
                continue

            # 处理标题
            if line.startswith('#'):
                if current_slide.title and not current_slide.bullets:
                    # 二级标题，新页面
                    slides.append(current_slide)
                    current_slide = SlideContent()
                title, level = self._parse_heading(line)
                if level == 1:
                    current_slide.title = title
                    current_slide.layout = LayoutType.COVER if not slides else LayoutType.TITLE
                else:
                    current_slide.title = title
                continue

            # 处理列表
            if self._is_list_item(line):
                bullet = self._parse_list_item(line)
                current_slide.bullets.append(bullet)
                current_slide.layout = LayoutType.CONTENT
                continue

            # 处理引用
            if line.startswith('>'):
                quote = self._parse_quote(line)
                if quote:
                    current_slide.subtitle = quote
                continue

            # 处理代码块
            if line.startswith('```'):
                if current_slide.code:
                    # 代码块结束
                    pass
                else:
                    current_slide.code_lang = self._parse_code_lang(line)
                    current_slide.layout = LayoutType.CODE
                continue

            # 处理表格
            if self._is_table_row(line):
                row = self._parse_table_row(line)
                if row:
                    current_slide.table_data.append(row)
                    current_slide.layout = LayoutType.TABLE
                continue

            # 处理图片
            if line.startswith('!'):
                image_match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
                if image_match:
                    current_slide.image_url = image_match.group(2)

        # 添加最后一页
        if current_slide.title or current_slide.bullets or current_slide.code:
            slides.append(current_slide)

        # 如果没有幻灯片，创建一个默认的
        if not slides:
            slides.append(SlideContent(
                layout=LayoutType.CONTENT,
                title="内容"
            ))

        return slides

    def _is_slide_separator(self, line: str) -> bool:
        """检查是否是幻灯片分隔符"""
        return line.strip() in ['---', '***', '___']

    def _is_comment(self, line: str) -> bool:
        """检查是否是注释指令"""
        return line.strip().startswith('<!--') and line.strip().endswith('-->')

    def _process_comment(self, line: str, slide: SlideContent):
        """处理注释指令"""
        content = line.strip()[4:-3].strip()

        if content.startswith('slide:'):
            parts = content.split(':', 1)
            if len(parts) == 2:
                key, value = parts
                if key == 'slide:layout':
                    try:
                        slide.layout = LayoutType(value.strip())
                    except ValueError:
                        pass
                elif key == 'slide:notes':
                    slide.notes = value.strip()
                elif key == 'slide:transition':
                    slide.transition = value.strip()
                elif key == 'slide:background':
                    slide.background = value.strip()

    def _parse_heading(self, line: str) -> tuple[str, int]:
        """解析标题"""
        match = re.match(r'^(#{1,6})\s+(.*)', line)
        if match:
            level = len(match.group(1))
            return match.group(2).strip(), level
        return line, 0

    def _is_list_item(self, line: str) -> bool:
        """检查是否是列表项"""
        return bool(re.match(r'^[\-\*\+]\s+', line)) or bool(re.match(r'^\d+\.\s+', line))

    def _parse_list_item(self, line: str) -> str:
        """解析列表项"""
        match = re.match(r'^[\-\*\+]\s+(.*)', line)
        if match:
            return match.group(1).strip()
        match = re.match(r'^\d+\.\s+(.*)', line)
        if match:
            return match.group(1).strip()
        return line.strip()

    def _parse_quote(self, line: str) -> str:
        """解析引用"""
        return line.lstrip('>').strip()

    def _parse_code_lang(self, line: str) -> str:
        """解析代码语言"""
        match = re.match(r'^```(\w*)', line)
        if match:
            return match.group(1)
        return ""

    def _is_table_row(self, line: str) -> bool:
        """检查是否是表格行"""
        return '|' in line and not line.strip().startswith('|---')

    def _parse_table_row(self, line: str) -> list[str]:
        """解析表格行"""
        cells = [c.strip() for c in line.split('|')]
        return [c for c in cells if c]


class SlidePlanner:
    """幻灯片规划器"""

    def __init__(self, slides: list[SlideContent], style_config: SlideStyleConfig):
        self.slides = slides
        self.style_config = style_config

    def plan(self) -> list[SlideContent]:
        """规划幻灯片结构"""
        for i, slide in enumerate(self.slides):
            # 自动推断布局
            if not slide.layout or slide.layout == LayoutType.CONTENT:
                slide.layout = self._infer_layout(slide)

            # 添加样式信息
            slide.primary_color = self.style_config.primary_color
            slide.secondary_color = self.style_config.secondary_color

        # 添加结尾页
        if self.slides and self.slides[-1].layout != LayoutType.ENDING:
            self.slides.append(SlideContent(
                layout=LayoutType.ENDING,
                title="谢谢",
                subtitle="欢迎提问"
            ))

        return self.slides

    def _infer_layout(self, slide: SlideContent) -> LayoutType:
        """推断最佳布局"""
        if slide.code:
            return LayoutType.CODE
        if slide.table_data:
            return LayoutType.TABLE
        if slide.image_url:
            return LayoutType.TWO_COLUMN
        if len(slide.bullets) <= 3:
            return LayoutType.CONTENT
        return LayoutType.CONTENT


class MarkdownToPPTX:
    """主转换器"""

    def __init__(self, input_path: str, output_path: Optional[str] = None,
                 style: str = "corporate", format: str = "ppt169",
                 title: Optional[str] = None):
        self.input_path = Path(input_path)
        self.output_path = self._resolve_output_path(output_path)
        self.style = SlideStyle(style)
        self.format = format
        self.title = title or self.input_path.stem
        self.style_config = SlideStyleConfig.from_style(self.style)

    def _resolve_output_path(self, output_path: Optional[str]) -> Path:
        """解析输出路径"""
        if output_path:
            return Path(output_path)

        # 自动生成输出路径
        return self.input_path.parent / f"{self.input_path.stem}_{self.style.value}.pptx"

    def convert(self) -> ConversionResult:
        """执行转换"""
        try:
            # 读取输入文件
            if not self.input_path.exists():
                return ConversionResult(
                    success=False,
                    error_message=f"文件不存在: {self.input_path}"
                )

            content = self.input_path.read_text(encoding='utf-8')

            # 解析 Markdown
            parser = MarkdownParser(content)
            slides = parser.parse()

            # 规划幻灯片
            planner = SlidePlanner(slides, self.style_config)
            slides = planner.plan()

            # 生成 PPTX
            self._generate_pptx(slides)

            # 收集统计信息
            stats = {
                "slide_count": len(slides),
                "input_file": str(self.input_path),
                "output_file": str(self.output_path),
                "style": self.style.value,
                "format": self.format
            }

            return ConversionResult(
                success=True,
                slides=slides,
                output_path=self.output_path,
                stats=stats
            )

        except Exception as e:
            return ConversionResult(
                success=False,
                error_message=f"转换失败: {str(e)}"
            )

    def _generate_pptx(self, slides: list[SlideContent]):
        """生成 PPTX 文件"""
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RgbColor
            from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
            from pptx.enum.shapes import MSO_SHAPE

            # 创建演示文稿
            prs = Presentation()
            prs.slide_width = Inches(13.333)  # 16:9
            prs.slide_height = Inches(7.5)

            for slide_content in slides:
                slide = prs.slides.add_slide(prs.slide_layouts[6])  # 空白布局

                # 添加内容
                self._add_slide_content(slide, slide_content, prs)

            # 保存
            prs.save(str(self.output_path))

        except ImportError:
            # 如果没有 python-pptx，生成 SVG 替代
            self._generate_svg_fallback(slides)

    def _add_slide_content(self, slide, content: SlideContent, prs):
        """添加幻灯片内容"""
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

        # 背景色
        if content.background:
            pass  # 需要额外处理

        # 添加标题
        if content.title:
            title_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(0.5),
                Inches(12), Inches(1)
            )
            tf = title_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = content.title
            p.font.size = Pt(44)
            p.font.bold = True
            p.font.color.rgb = RgbColor(0x1E, 0x3A, 0x5F)

        # 添加要点
        if content.bullets:
            y_pos = 2.0 if content.title else 0.5
            content_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(y_pos),
                Inches(12), Inches(5)
            )
            tf = content_box.text_frame
            tf.word_wrap = True

            for i, bullet in enumerate(content.bullets):
                if i == 0:
                    p = tf.paragraphs[0]
                else:
                    p = tf.add_paragraph()
                p.text = f"• {bullet}"
                p.font.size = Pt(24)
                p.space_after = Pt(12)

        # 添加代码
        if content.code:
            code_box = slide.shapes.add_textbox(
                Inches(0.5), Inches(2.5),
                Inches(12), Inches(4)
            )
            tf = code_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = content.code
            p.font.size = Pt(14)
            p.font.name = "Consolas"
            p.font.color.rgb = RgbColor(0x2D, 0x2D, 0x2D)

        # 添加备注
        if content.notes:
            notes_slide = slide.notes_slide
            notes_tf = notes_slide.notes_text_frame
            notes_tf.text = content.notes

    def _generate_svg_fallback(self, slides: list[SlideContent]):
        """SVG 降级方案"""
        output_dir = self.output_path.parent / f"{self.output_path.stem}_svg"
        output_dir.mkdir(exist_ok=True)

        for i, slide in enumerate(slides, 1):
            svg_content = self._render_slide_to_svg(slide, i)
            svg_path = output_dir / f"slide_{i:02d}.svg"
            svg_path.write_text(svg_content, encoding='utf-8')

        # 更新输出路径为目录
        self.output_path = output_dir

    def _render_slide_to_svg(self, content: SlideContent, index: int) -> str:
        """渲染幻灯片为 SVG"""
        view_box = "0 0 1280 720"
        title = self._escape_xml(content.title)
        bullets = "\n".join([
            f'<text x="60" y="{200 + i * 40}" font-size="24">{self._escape_xml(b)}</text>'
            for i, b in enumerate(content.bullets)
        ])

        return f'''<?xml version="1.0" encoding="UTF-8"?>
<svg xmlns="http://www.w3.org/2000/svg" viewBox="{view_box}">
    <rect width="1280" height="720" fill="{content.background or self.style_config.background_color}"/>
    <text x="60" y="80" font-size="48" font-weight="bold" fill="{self.style_config.primary_color}">{title}</text>
    {bullets}
</svg>'''

    def _escape_xml(self, text: str) -> str:
        """转义 XML 特殊字符"""
        return (text
                .replace('&', '&amp;')
                .replace('<', '&lt;')
                .replace('>', '&gt;')
                .replace('"', '&quot;'))


def main():
    """命令行入口"""
    parser = argparse.ArgumentParser(
        description='将 Markdown 文件转换为 PowerPoint 演示文稿',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
示例:
  python md2pptx.py input.md
  python md2pptx.py input.md --style corporate
  python md2pptx.py input.md --output presentation.pptx
  python md2pptx.py input.md --title "演示标题"
        """
    )

    parser.add_argument('input', help='Markdown 文件路径')
    parser.add_argument('-o', '--output', help='输出 PPTX 文件路径')
    parser.add_argument('-s', '--style', default='corporate',
                        choices=['corporate', 'minimal', 'academic', 'tech', 'dark', 'creative'],
                        help='演示风格')
    parser.add_argument('-f', '--format', default='ppt169',
                        choices=['ppt169', 'ppt43'],
                        help='幻灯片比例')
    parser.add_argument('-t', '--title', help='演示文稿标题')
    parser.add_argument('--verbose', '-v', action='store_true',
                        help='详细输出')
    parser.add_argument('--stats', action='store_true',
                        help='输出统计信息')

    args = parser.parse_args()

    # 执行转换
    converter = MarkdownToPPTX(
        input_path=args.input,
        output_path=args.output,
        style=args.style,
        format=args.format,
        title=args.title
    )

    result = converter.convert()

    if result.success:
        print(f"✅ 转换成功!")
        print(f"📄 输出文件: {result.output_path}")

        if args.stats:
            print(f"\n📊 统计信息:")
            for key, value in result.stats.items():
                print(f"   {key}: {value}")
    else:
        print(f"❌ 转换失败: {result.error_message}")
        sys.exit(1)


if __name__ == '__main__':
    main()
