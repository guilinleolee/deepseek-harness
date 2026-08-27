#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Word模板生成器 - 创建招投标文档样式模板
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

def create_bidding_template():
    """创建招投标文档Word模板"""
    doc = Document()

    # 设置页面边距
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    # 创建样式
    styles = doc.styles

    # 标题样式 - 黑体二号
    title_style = styles.add_style('BiddingTitle', WD_STYLE_TYPE.PARAGRAPH)
    title_style.font.name = '黑体'
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_style.paragraph_format.space_before = Pt(0)
    title_style.paragraph_format.space_after = Pt(12)

    # 一级标题样式 - 黑体三号
    h1_style = styles.add_style('BiddingH1', WD_STYLE_TYPE.PARAGRAPH)
    h1_style.font.name = '黑体'
    h1_style.font.size = Pt(16)
    h1_style.font.bold = True
    h1_style.paragraph_format.space_before = Pt(12)
    h1_style.paragraph_format.space_after = Pt(6)

    # 二级标题样式 - 黑体四号
    h2_style = styles.add_style('BiddingH2', WD_STYLE_TYPE.PARAGRAPH)
    h2_style.font.name = '黑体'
    h2_style.font.size = Pt(14)
    h2_style.font.bold = True
    h2_style.paragraph_format.space_before = Pt(6)
    h2_style.paragraph_format.space_after = Pt(3)

    # 三级标题样式 - 黑体小四
    h3_style = styles.add_style('BiddingH3', WD_STYLE_TYPE.PARAGRAPH)
    h3_style.font.name = '黑体'
    h3_style.font.size = Pt(12)
    h3_style.font.bold = True
    h3_style.paragraph_format.space_before = Pt(3)
    h3_style.paragraph_format.space_after = Pt(3)

    # 正文样式 - 宋体小四
    body_style = styles.add_style('BiddingBody', WD_STYLE_TYPE.PARAGRAPH)
    body_style.font.name = '宋体'
    body_style.font.size = Pt(12)
    body_style.paragraph_format.first_line_indent = Cm(0.74)  # 两个字符
    body_style.paragraph_format.line_spacing = 1.5
    body_style.paragraph_format.space_before = Pt(0)
    body_style.paragraph_format.space_after = Pt(0)

    # 列表样式 - 宋体小四
    list_style = styles.add_style('BiddingList', WD_STYLE_TYPE.PARAGRAPH)
    list_style.font.name = '宋体'
    list_style.font.size = Pt(12)
    list_style.paragraph_format.left_indent = Cm(0.74)
    list_style.paragraph_format.line_spacing = 1.5

    # 表格标题样式
    table_title_style = styles.add_style('BiddingTableTitle', WD_STYLE_TYPE.PARAGRAPH)
    table_title_style.font.name = '黑体'
    table_title_style.font.size = Pt(12)
    table_title_style.font.bold = True
    table_title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    table_title_style.paragraph_format.space_before = Pt(6)
    table_title_style.paragraph_format.space_after = Pt(3)

    # 页眉页脚样式
    header_style = styles.add_style('BiddingHeader', WD_STYLE_TYPE.PARAGRAPH)
    header_style.font.name = '宋体'
    header_style.font.size = Pt(9)
    header_style.font.color.rgb = RGBColor(128, 128, 128)
    header_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 封面样式
    cover_title_style = styles.add_style('CoverTitle', WD_STYLE_TYPE.PARAGRAPH)
    cover_title_style.font.name = '黑体'
    cover_title_style.font.size = Pt(36)
    cover_title_style.font.bold = True
    cover_title_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    cover_subtitle_style = styles.add_style('CoverSubtitle', WD_STYLE_TYPE.PARAGRAPH)
    cover_subtitle_style.font.name = '黑体'
    cover_subtitle_style.font.size = Pt(18)
    cover_subtitle_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 创建示例封面
    doc.add_paragraph('招投标文件', style='CoverTitle')
    doc.add_paragraph()
    doc.add_paragraph('项目名称：_______________', style='CoverSubtitle')
    doc.add_paragraph('项目编号：_______________', style='CoverSubtitle')
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph('招标单位：_______________', style='CoverSubtitle')
    doc.add_paragraph('投标单位：_______________', style='CoverSubtitle')
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph('日期：______年______月______日', style='CoverSubtitle')
    doc.add_page_break()

    # 创建示例目录
    doc.add_paragraph('目  录', style='BiddingTitle')
    doc.add_paragraph()
    toc_items = [
        '第一章 投标须知.........................................................1',
        '第二章 项目概况.........................................................3',
        '第三章 技术方案.........................................................5',
        '第四章 报价明细.........................................................15',
        '第五章 项目团队.........................................................18',
        '第六章 服务承诺.........................................................20',
        '附件.......................................................................22',
    ]
    for item in toc_items:
        doc.add_paragraph(item, style='BiddingBody')
    doc.add_page_break()

    # 创建示例正文
    doc.add_paragraph('第一章 投标须知', style='BiddingTitle')
    doc.add_paragraph()

    doc.add_paragraph('一、项目概述', style='BiddingH1')
    doc.add_paragraph('本项目旨在[项目目标描述]，欢迎符合条件的投标方参与投标。', style='BiddingBody')

    doc.add_paragraph('二、投标资格要求', style='BiddingH1')
    doc.add_paragraph('1. 投标方应具备相应的资质，在相关行业有一定年限的从业经验。', style='BiddingList')
    doc.add_paragraph('2. 提供所需的资质证明文件，包括但不限于营业执照、行业资质证书等。', style='BiddingList')
    doc.add_paragraph('3. 具有良好的商业信誉和健全的财务会计制度。', style='BiddingList')

    doc.add_paragraph('三、投标文件要求', style='BiddingH1')
    doc.add_paragraph('投标文件应包含以下内容：', style='BiddingBody')
    doc.add_paragraph('（一）商务部分', style='BiddingH2')
    doc.add_paragraph('1. 投标函', style='BiddingList')
    doc.add_paragraph('2. 法定代表人授权书', style='BiddingList')
    doc.add_paragraph('3. 营业执照副本复印件', style='BiddingList')

    doc.add_paragraph('（二）技术部分', style='BiddingH2')
    doc.add_paragraph('1. 技术方案', style='BiddingList')
    doc.add_paragraph('2. 实施计划', style='BiddingList')
    doc.add_paragraph('3. 质量保证措施', style='BiddingList')

    doc.add_paragraph('（三）报价部分', style='BiddingH2')
    doc.add_paragraph('1. 投标报价汇总表', style='BiddingList')
    doc.add_paragraph('2. 分项报价明细表', style='BiddingList')

    # 创建示例表格
    doc.add_paragraph()
    doc.add_paragraph('表1-1 投标报价汇总表', style='BiddingTableTitle')

    table = doc.add_table(rows=5, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 设置表头
    header_cells = table.rows[0].cells
    headers = ['序号', '项目名称', '金额（万元）', '备注']
    for i, header in enumerate(headers):
        header_cells[i].text = header
        for paragraph in header_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True

    # 填充示例数据
    data = [
        ['1', '人力资源费用', '_______', ''],
        ['2', '硬件设备费用', '_______', ''],
        ['3', '软件授权费用', '_______', ''],
        ['合计', '', '_______', ''],
    ]
    for i, row_data in enumerate(data):
        row = table.rows[i + 1]
        for j, cell_data in enumerate(row_data):
            row.cells[j].text = cell_data

    # 保存模板
    template_path = os.path.join(os.path.dirname(__file__), '..', 'templates', 'docx', 'bidding_template.docx')
    os.makedirs(os.path.dirname(template_path), exist_ok=True)
    doc.save(template_path)
    print(f"模板已保存到: {template_path}")
    return template_path


def create_proposal_template():
    """创建投标文件专用模板"""
    doc = Document()

    # 设置页面边距
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    # 封面
    doc.add_paragraph()
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('投标文件')
    run.font.name = '黑体'
    run.font.size = Pt(44)
    run.font.bold = True

    doc.add_paragraph()
    doc.add_paragraph()

    # 项目信息
    info_items = [
        '项目名称：________________________',
        '项目编号：________________________',
        '',
        '投标单位：________________________',
        '法定代表人：______________________',
        '联系人：__________________________',
        '联系电话：________________________',
        '',
        '投标日期：______年______月______日',
    ]
    for item in info_items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if item:
            run = p.add_run(item)
            run.font.name = '宋体'
            run.font.size = Pt(14)

    doc.add_page_break()

    # 投标函
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('投标函')
    run.font.name = '黑体'
    run.font.size = Pt(22)
    run.font.bold = True

    doc.add_paragraph()
    doc.add_paragraph('致：________________________')
    doc.add_paragraph()
    doc.add_paragraph('    我方已仔细阅读并充分理解________________________项目（项目编号：________________________）招标文件的全部内容，现决定参加该项目投标。')
    doc.add_paragraph()
    doc.add_paragraph('    我方承诺：')
    doc.add_paragraph('    一、所提供的所有资料真实、有效，并愿意承担相应的法律责任；')
    doc.add_paragraph('    二、如我方中标，将严格按照招标文件和投标文件的约定履行合同义务；')
    doc.add_paragraph('    三、投标有效期为______天，在此期间内我方承诺不撤回投标。')
    doc.add_paragraph()
    doc.add_paragraph('    投标报价：人民币________万元（大写：________________________）')
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph('投标单位（盖章）：________________________')
    doc.add_paragraph('法定代表人或授权代表（签字）：________________________')
    doc.add_paragraph('日期：______年______月______日')

    # 保存模板
    template_path = os.path.join(os.path.dirname(__file__), '..', 'templates', 'docx', 'proposal_template.docx')
    os.makedirs(os.path.dirname(template_path), exist_ok=True)
    doc.save(template_path)
    print(f"投标文件模板已保存到: {template_path}")
    return template_path


if __name__ == '__main__':
    print("正在创建Word样式模板...")
    create_bidding_template()
    create_proposal_template()
    print("模板创建完成！")